from pathlib import Path
from zipfile import ZipFile, ZIP_DEFLATED
from lxml import etree as E
from io import BytesIO
from openpyxl import load_workbook
import json,copy,hashlib
root=Path(__file__).resolve().parent.parent
src=root.parent/'content-rewards-business-v12/output/Wringy-Content-Rewards-Business-Plan.pptx'
n={'c':'http://schemas.openxmlformats.org/drawingml/2006/chart'}
with ZipFile(src) as old, ZipFile(root/'build/draft.pptx') as new:
 replacements={k:new.read(k) for k in ['ppt/slides/slide13.xml','ppt/slides/_rels/slide13.xml.rels','ppt/notesSlides/notesSlide13.xml']}
 chart=E.fromstring(old.read('ppt/slides/charts/chart4.xml')); nc=E.fromstring(new.read('ppt/slides/charts/chart4.xml'))
 ser=chart.find('.//c:ser',n)
 for dp in list(ser.findall('c:dPt',n)):ser.remove(dp)
 idx=list(ser).index(ser.find('c:cat',n))
 for dp in nc.findall('.//c:ser/c:dPt',n):ser.insert(idx,copy.deepcopy(dp));idx+=1
 for tag,kind,values in [('cat','str',['研发','储备资金']),('val','num',[50000,20000])]:
  cache=ser.find(f'c:{tag}/c:{kind}Ref/c:{kind}Cache',n)
  for pt in list(cache.findall('c:pt',n)):cache.remove(pt)
  cache.find('c:ptCount',n).set('val','2')
  for i,v in enumerate(values):
   pt=E.SubElement(cache,'{'+n['c']+'}pt',idx=str(i));E.SubElement(pt,'{'+n['c']+'}v').text=str(v)
  f=ser.find(f'c:{tag}/c:{kind}Ref/c:f',n);f.text=f.text.replace('$4','$3')
 replacements['ppt/slides/charts/chart4.xml']=E.tostring(chart,xml_declaration=True,encoding='UTF-8')
 wb=load_workbook(BytesIO(old.read('ppt/embeddings/chart-data-snapshot-004.xlsx')));ws=wb.active;ws['A3']='储备资金';ws['B3']=20000;ws.delete_rows(4)
 out=BytesIO();wb.save(out);replacements['ppt/embeddings/chart-data-snapshot-004.xlsx']=out.getvalue()
 with ZipFile(root/'build/candidate.pptx','w',ZIP_DEFLATED) as z:
  for item in old.infolist():z.writestr(item,replacements.get(item.filename,old.read(item.filename)))
 with ZipFile(root/'build/candidate.pptx') as z:
  changed=[k for k in old.namelist() if z.read(k)!=old.read(k)]
  assert sorted(changed)==sorted(replacements)
 print('Changed parts:',changed)
 (root/'build/changed-parts.json').write_text(json.dumps(changed,indent=2))
